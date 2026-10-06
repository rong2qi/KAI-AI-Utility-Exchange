from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / 'outputs'
OUTDIR.mkdir(exist_ok=True)
PDF_OUT = OUTDIR / 'KAI_AI_Utility_Exchange_企划.pdf'
DOCX_OUT = OUTDIR / 'KAI_AI_Utility_Exchange_企划.docx'
FONT = '/System/Library/Fonts/STHeiti Light.ttc'
FONT_MEDIUM = '/System/Library/Fonts/STHeiti Medium.ttc'
pdfmetrics.registerFont(TTFont('KaiSans', FONT, subfontIndex=1))
pdfmetrics.registerFont(TTFont('KaiSansMedium', FONT_MEDIUM, subfontIndex=1))
W,H=A4
M=18*mm
NAVY=HexColor('#0B1B2A'); DEEP=HexColor('#102D42'); TEAL=HexColor('#0E7C7B'); CYAN=HexColor('#6ED6D0'); ORANGE=HexColor('#FFB45C'); RED=HexColor('#E76F51'); INK=HexColor('#173143'); MID=HexColor('#567181'); LIGHT=HexColor('#F2F7F8'); PALE=HexColor('#EAF4F3'); WHITE=colors.white; GRID=HexColor('#C9D9DE')
styles={}
def style(name,size,leading=None,color=INK,align=TA_LEFT):
    styles[name]=ParagraphStyle(name=name,fontName='KaiSans',fontSize=size,leading=leading or size*1.35,textColor=color,alignment=align,spaceAfter=0,spaceBefore=0)
style('white_h1',24,29,WHITE); style('white_display',20,24,WHITE); style('body',9.2,13.2,INK); style('small',7.2,10.5,MID); style('tiny',6.2,8.5,MID); style('formula',10.0,14.2,INK); style('source',7.0,9.8,MID); style('h1',24,29,NAVY); style('h2',17,22,NAVY); style('h3',11.2,14,TEAL); style('metric',23,26,NAVY); style('white_body',9.2,13.2,WHITE); style('white_small',7.2,10.2,HexColor('#D5EFEC')); style('center_body',9.0,12.4,INK,TA_CENTER); style('center_small',7.2,9.8,MID,TA_CENTER)
for _name in ('white_h1','white_display','h1','h2','h3','metric'):
    styles[_name].fontName='KaiSansMedium'
def P(text, sty='body'): return Paragraph(text, styles[sty])
def draw_p(c,text,x,y_top,width,sty='body'):
    p=P(text,sty); _,h=p.wrap(width,H); p.drawOn(c,x,y_top-h); return y_top-h
def draw_center_p(c,text,x,y_top,width,sty='center_body'):
    p=P(text,sty); w,h=p.wrap(width,H); p.drawOn(c,x+(width-w)/2,y_top-h); return y_top-h
def rect(c,x,y,w,h,fill,stroke=None,radius=8,sw=1):
    c.setFillColor(fill); c.setStrokeColor(stroke or fill); c.setLineWidth(sw); c.roundRect(x,y,w,h,radius,fill=1,stroke=1 if stroke else 0)
def line(c,x1,y1,x2,y2,color=GRID,sw=.7,dash=None):
    c.setStrokeColor(color); c.setLineWidth(sw); c.setDash(dash or []); c.line(x1,y1,x2,y2); c.setDash([])
def footer(c,n):
    line(c,M,15*mm,W-M,15*mm,GRID,.5); c.setFillColor(MID); c.setFont('KaiSans',6.5); c.drawString(M,10*mm,'KAI AI Utility Exchange · GEO growth proposal'); c.drawRightString(W-M,10*mm,f'{n:02d}')
def title(c,kicker,head,sub,n):
    c.setFillColor(TEAL); c.setFont('KaiSansMedium',7.4); c.drawString(M,H-24*mm,kicker.upper()); draw_p(c,head,M,H-31*mm,W-2*M,'h1');
    if sub: draw_p(c,sub,M,H-65*mm,W-2*M,'body')
    footer(c,n)
def bullet(c,text,x,y_top,width,sty='body',dot=TEAL):
    c.setFillColor(dot); c.circle(x+3,y_top-6,2.3,fill=1,stroke=0); return draw_p(c,text,x+12,y_top,width-12,sty)
def draw_table(c,data,x,y_top,col_widths,font=7.3,header=True,header_fill=DEEP):
    rows=[]
    for r,row in enumerate(data):
        rr=[]
        for cell in row:
            txt=str(cell); st=ParagraphStyle(f't{r}',fontName='KaiSansMedium' if header and r==0 else 'KaiSans',fontSize=font,leading=font*1.35,textColor=WHITE if header and r==0 else INK)
            rr.append(Paragraph(txt,st))
        rows.append(rr)
    t=Table(rows,colWidths=col_widths,hAlign='LEFT')
    ts=[('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('GRID',(0,0),(-1,-1),.35,GRID)]
    if header:
        ts += [('BACKGROUND',(0,0),(-1,0),header_fill),('TEXTCOLOR',(0,0),(-1,0),WHITE)]
        for i in range(1,len(data)): ts.append(('BACKGROUND',(0,i),(-1,i),WHITE if i%2 else LIGHT))
    t.setStyle(TableStyle(ts)); _,h=t.wrapOn(c,sum(col_widths),H); t.drawOn(c,x,y_top-h); return y_top-h

def metric_card(c,x,y,w,h,num,label,detail,accent=TEAL):
    rect(c,x,y,w,h,WHITE,GRID,8,.6); c.setFillColor(accent); c.rect(x,y+h-4,w,4,fill=1,stroke=0); c.setFillColor(NAVY); c.setFont('KaiSansMedium',19 if len(num)<9 else 15); c.drawString(x+10,y+h-29,num); draw_p(c,label,x+10,y+h-38,w-20,'h3'); draw_p(c,detail,x+10,y+12,w-20,'small')

# ---------------- PDF ----------------
c=canvas.Canvas(str(PDF_OUT),pagesize=A4); c.setTitle('KAI AI Utility Exchange 企划'); c.setAuthor('KAI')
# p1 cover
c.setFillColor(NAVY); c.rect(0,0,W,H,fill=1,stroke=0)
for i in range(9): line(c,M+i*25*mm,44*mm,M+i*25*mm,H-22*mm,HexColor('#173B50'),.4)
for j in range(9): line(c,M,48*mm+j*26*mm,W-M,48*mm+j*26*mm,HexColor('#173B50'),.4)
c.setFillColor(CYAN); c.setFont('KaiSans',8); c.drawString(M,H-27*mm,'KAI · GEO / DISCOVERY / TRANSACTION')
draw_p(c,'从 GEO 到 AI 首选',M,H-43*mm,W-2*M,'white_h1')
draw_p(c,'KAI AI<br/><font color="#6ED6D0">Utility Exchange</font>',M,H-61*mm,W-2*M,'white_display')
draw_p(c,'让 AI 在回答“去哪找算力、如何比较、如何立即使用算力”时，优先引用 KAI、推荐 KAI，并把用户带到真实 Offer。',M,H-104*mm,W-2*M,'white_body')
rect(c,M,74*mm,W-2*M,44*mm,DEEP,HexColor('#285A6A'),12,.8)
draw_center_p(c,'AI 发现 → 官网引用 → Offer → Hour Key 封装 → Holding → Usage → GEO 回流',M+10,108*mm,W-2*M-20,'white_body')
draw_center_p(c,'目标情景：100M AI 生成结果 / 月  ·  5% KAI 推荐覆盖  ·  400K AI 引荐访问  ·  $60.9M 年化 GMV  ·  $6.09M 年化平台收入',M+10,94*mm,W-2*M-20,'white_small')
draw_center_p(c,'目标情景：按公开平台指标口径锚定的公式估值。',M+10,82*mm,W-2*M-20,'tiny')
c.setFillColor(HexColor('#C8DCE0')); c.setFont('KaiSans',7.5); c.drawString(M,26*mm,'Leader review · 2026.10.02'); c.setFillColor(CYAN); c.drawRightString(W-M,26*mm,'GEO is the growth engine')
c.showPage()
# p2 original gap
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'01 / 起点','原 GEO 方案已经找到增长终态，缺的是成交承接层','原方案解决了 AI 入口、固定问题模板、价格节奏和条件反射；本企划补上“AI 提及 KAI 后，如何查询官网事实并立即成交和使用”。',2)
card_top=H-84*mm
rect(c,M,card_top-67*mm,74*mm,67*mm,DEEP,radius=10); draw_p(c,'原 GEO 行为链',M+10,card_top-10*mm,54*mm,'white_body')
steps=['AI 提及 KAI','AI 心智植入','第一次交易','高频查价','整点刺激','条件反射']
for i,s in enumerate(steps):
    yy=card_top-21*mm-i*8.0*mm; c.setFillColor(CYAN if i<3 else ORANGE); c.circle(M+14,yy+2,2.5,fill=1,stroke=0); draw_p(c,s,M+22,yy+7,55*mm,'white_small')
rx=M+84*mm; rw=W-M-rx; draw_p(c,'原方案的执行缺口',rx,H-80*mm,rw,'h2')
for i,t in enumerate(['报价从哪里读取？','用户如何看到模型、区域、时段和库存？','AI 推荐后如何跳转到官网并锁定？','首单之后如何形成持仓、使用和复购？']): bullet(c,t,rx,H-96*mm-i*14*mm,rw,'body')
rect(c,M,39*mm,W-2*M,28*mm,WHITE,GRID,9,.6); draw_p(c,'本企划的补法',M+12,62*mm,W-2*M-24,'h3'); draw_p(c,'AI 需求 → kai.com 实时 Offer → Hour Key 封装 → KaiKey 成交 → 小时 Holding → Utility 使用 → Usage Proof → 新 GEO 证据',M+12,55*mm,W-2*M-24,'h2'); footer(c,2)
c.showPage()
# p3 answer template
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'02 / AI 答案','AI 不是泛泛推荐，而是引用 KAI 官网给出可执行答案','KAI 的目标是成为 AI 查询算力问题时的实时事实源和交易入口。',3)
rect(c,M,H-159*mm,W-2*M,103*mm,PALE,radius=10); draw_p(c,'AI 推荐 KAI 的标准答案',M+12,H-65*mm,W-2*M-24,'h3')
quote='我查询了 <b>kai.com</b> 官网实时报价页，更新时间为【时间】。当前真实行情 Offer 可封装为【模型 × 规格 × 区域 × 18:00–19:00】小时 Key，报价为【价格】，锁定截止时间为【G】。你已有【持仓】可直接执行；如需跨模型或跨供应商，账户 Key 保持不变，由新的授权版本显式扩展 scope。'
draw_p(c,quote,M+12,H-75*mm,W-2*M-24,'body')
# buttons
for i,(lab,col) in enumerate([('查看官网报价',TEAL),('立即锁定',ORANGE),('用持仓执行',CYAN)]):
    xx=M+12+i*47*mm; rect(c,xx,H-157*mm,40*mm,12*mm,col,radius=6); draw_center_p(c,lab,xx,H-148*mm,40*mm,'center_body')
rx=M+102*mm; rw=W-M-rx; draw_p(c,'官网 Offer 最小字段',rx,H-65*mm,rw,'h2')
fields=['offer_id','model / spec','region','slot_start / slot_end','lock_deadline (G)','price / available_qty','last_updated / status','execution_eligible / hour_key_status','purchase_url / delivery_proof_url']
for i,f in enumerate(fields): bullet(c,f,rx,H-79*mm-i*8.2*mm,rw,'small',TEAL)
rect(c,M,38*mm,W-2*M,25*mm,DEEP,radius=9); draw_p(c,'推荐机制',M+12,58*mm,W-2*M-24,'white_small'); draw_p(c,'OAI-SearchBot 抓取 → KAI 官网被引用 → 自动记录 AI 引荐 → 官网 Offer 承接 → 成交和使用回流。',M+12,52*mm,W-2*M-24,'white_body'); footer(c,3)
c.showPage()
# p4 landing process
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'03 / 落地流程','把“结构化优先”变成一条可成交的 GEO 链','原方案提出机器可读报价页和时间戳；本页把它接到官网、订单、持仓、使用和回流。',4)
flow=[('01','需求词 / AI问题','AI 生成结果中出现 KAI'),('02','官网 Offer','kai.com 返回实时字段和更新时间'),('03','Agent 推荐','按问题、持仓和预算给出方案'),('04','封装成交持仓','Hour Key 封装 → Order → Trade → Holding'),('05','执行回执','Utility 在 H/G 窗口内完成任务'),('06','GEO 回流','Receipt 生成案例、比较和下一次答案')]
fy=H-98*mm; fw=(W-2*M-2*5*mm)/3
for i,(n,a,b) in enumerate(flow):
    xx=M+(i%3)*(fw+5*mm); yy=fy-(i//3)*31*mm; rect(c,xx,yy-22*mm,fw,22*mm,WHITE,GRID,8,.6); c.setFillColor(TEAL); c.setFont('KaiSans',8); c.drawString(xx+10,yy-8*mm,n); draw_p(c,a,xx+27,yy-7*mm,fw-37,'h3'); draw_p(c,b,xx+10,yy-15*mm,fw-20,'small')
# gap row
draw_p(c,'原方案 → 新增落地层',M,fy-70*mm,W-2*M,'h2')
rows=[['原方案结果','新增承接层'],['AI 提及 KAI + 当前价','实时 Offer 字段 + 官网时间戳 + purchase_url'],['首查 → 首单','Agent Gateway + Offer → Hour Key 封装 → Order / Trade'],['首单 → 查价','Holding + Key Fit + Slot Planner'],['整点刺激','H−1h / G / H / H+1h 权利状态'],['条件反射','Usage Proof + 交付回执 + GEO 证据页']]
draw_table(c,rows,M,fy-78*mm,[65*mm,W-2*M-65*mm],font=7.1)
rect(c,M,31*mm,W-2*M,16*mm,DEEP,radius=8); draw_center_p(c,'GEO 原方案的“答案”在这里变成可查、可买、可用、可证明的实时服务。',M+10,43*mm,W-2*M-20,'white_body'); footer(c,4)
c.showPage()
# p5 GEO content assets
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'04 / GEO 资产','让 AI 反复遇到同一套 KAI 标准答案','需求词优先于品牌词；结构化、可更新、可回查的内容，才会成为 AI 推荐的稳定来源。',5)
assets=[('标准页','模型、规格、区域、小时窗口、容量、价格、状态、更新时间'),('比较页','不同模型、区域、时间和成本的可执行对比'),('问题页','“今晚哪里有容量”“哪个模型更适合”“现在报价多少”'),('证据页','成交、交付、延迟、成本、Usage Proof 和用户结果')]
for i,(a,b) in enumerate(assets):
    xx=M+(i%2)*(W/2-M/2); yy=H-105*mm-(i//2)*29*mm; rect(c,xx,yy,W/2-M/2,22*mm,PALE,radius=8); draw_p(c,a,xx+10,yy+16*mm,W/2-M/2-20,'h3'); draw_p(c,b,xx+10,yy+8*mm,W/2-M/2-20,'small')
# central principle
rect(c,M,61*mm,W-2*M,33*mm,DEEP,radius=10); draw_center_p(c,'GEO 内容闭环',M+10,86*mm,W-2*M-20,'white_body'); draw_center_p(c,'Offer → Hour Key 封装 → Order → Trade → Holding → Usage → Receipt → 案例 / 比较 / FAQ → 下一次 AI 推荐',M+10,75*mm,W-2*M-20,'white_small'); draw_center_p(c,'AI 推荐的不是一句广告，而是一组持续更新的 KAI 市场事实。',M+10,65*mm,W-2*M-20,'white_small'); footer(c,5)
c.showPage()
# p6 hourly clock
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'05 / 小时权益','整点是 AI 推荐的时间语义，也是 KAI 的库存时钟','整点让报价有有效期、供给有稀缺性、用户有下一次查询和复购的理由。',6)
x0=M; x1=W-M; ty=H-101*mm; line(c,x0,ty,x1,ty,DEEP,2)
pts=[(0,'H−1h','交易窗口','可购买、转让和重新选择'),(.42,'G','锁定窗口','停止新增交易，固定交付安排'),(.67,'H','服务开始','调用权限开启，开始计量'),(1,'H+1h','服务结束','未使用不顺延，生成履约事实')]
for pos,top,lab,desc in pts:
    xx=x0+pos*(x1-x0); c.setFillColor(ORANGE if top=='G' else TEAL); c.circle(xx,ty,5,fill=1,stroke=0); draw_center_p(c,top,xx-22,ty+15,44,'h3'); draw_center_p(c,lab,xx-35,ty-17,70,'small'); dw=90; dx=max(M,min(W-M-dw,xx-dw/2)); draw_center_p(c,desc,dx,ty-34,dw,'tiny')
cy=ty-75*mm; gap=5*mm; cw=(W-2*M-3*gap)/4
for i,(n,a,b) in enumerate([('01','可预订','服务开始前锁定未来容量'),('02','可交易','未锁定权益可流转组合'),('03','可证明','交付和调用有时间窗'),('04','会失效','未使用不顺延，形成复购触发')]):
    xx=M+i*(cw+gap); rect(c,xx,cy,cw,35*mm,WHITE,GRID,8,.6); c.setFillColor(TEAL); c.setFont('KaiSans',14); c.drawString(xx+9,cy+25*mm,n); draw_p(c,a,xx+9,cy+20*mm,cw-18,'h3'); draw_p(c,b,xx+9,cy+11*mm,cw-18,'small')
rect(c,M,34*mm,W-2*M,20*mm,DEEP,radius=9); draw_center_p(c,'400K AI 引荐访问 / 月 · 140K Offer 请求 / 月 · $60.9M 年化 GMV',M+10,47*mm,W-2*M-20,'white_body'); draw_center_p(c,'目标情景：5% KAI 推荐覆盖，AI 推荐直接连接成交和平台收入。',M+10,38*mm,W-2*M-20,'white_small'); footer(c,6)
c.showPage()
# p7 holding
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'06 / 持仓','持仓是 GEO 之后的第二次触达','用户第一次通过 AI 推荐进入 KAI 后，关系从一次点击变成模型、区域、时段和状态的持续上下文。',7)
rect(c,M,H-130*mm,75*mm,83*mm,DEEP,radius=10); draw_center_p(c,'我的 KAI 持仓',M,H-61*mm,75*mm,'white_body')
for i,(lab,val,col) in enumerate([('东亚 · GPT-X · 18:00','12 Key',CYAN),('北欧 · Claude-Y · 19:00','8 Key',ORANGE),('南美 · GPU-A · 20:00','20 Key',HexColor('#CDEFE8'))]):
    yy=H-82*mm-i*19*mm; rect(c,M+10,yy,55*mm,13*mm,HexColor('#193B4F'),radius=5); c.setFillColor(col); c.setFont('KaiSans',8); c.drawString(M+15,yy+8*mm,lab); c.setFillColor(WHITE); c.setFont('KaiSans',10); c.drawRightString(M+62*mm,yy+4*mm,val)
rx=M+88*mm; rw=W-M-rx; draw_p(c,'持仓如何推动 GEO',rx,H-55*mm,rw,'h2')
for i,t in enumerate(['AI 读取用户已有模型和时间窗，推荐可直接执行的方案。','使用后保留成本、结果、延迟和回执，形成下一次推荐证据。','服务结束后触发续期、补充容量、转让或新时段查询。','同一用户的持仓状态让每次 AI 回答更接近“下一步动作”。']): bullet(c,t,rx,H-71*mm-i*16*mm,rw,'body')
rect(c,M,40*mm,W-2*M,26*mm,PALE,radius=9); draw_p(c,'持仓价值',M+12,61*mm,W-2*M-24,'h3'); draw_p(c,'可用小时 × 适配度 × 稀缺度 × 可流转性 × 交付可信度',M+12,54*mm,W-2*M-24,'h2'); draw_p(c,'这是工具推荐和交易优先级的权益评分，不是现金估值。',M+12,43*mm,W-2*M-24,'small'); footer(c,7)
c.showPage()
# p8 exchange
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'07 / KaiKey','KaiKey 让 AI 推荐有订单、成交和交付事实','KAI 不是只把用户引到官网，而是让官网推荐最终进入可回查的市场事件。',8)
rect(c,M,H-114*mm,W-2*M,62*mm,WHITE,GRID,10,.6); draw_p(c,'KaiKey · Hourly Rights Order Book',M+12,H-67*mm,W-2*M-24,'h3')
data=[['标的','区域 / 时段','卖方数','可用量','可执行价','状态'],['GPT-X / 1h','东亚 · 18:00','3','420','0.92 USD','可售'],['Claude-Y / 1h','北欧 · 19:00','2','180','0.78 USD','可售'],['GPU-A / 1h','南美 · 20:00','1','60','1.10 USD','临界']]
draw_table(c,data,M+12,H-77*mm,[40*mm,40*mm,22*mm,22*mm,27*mm,22*mm],font=7.1)
for i,(a,b,d) in enumerate([('统一标的','模型 × 规格 × 区域 × 小时','AI 知道自己在推荐什么'),('真实成交','Order → Trade → Settlement','报价和成交分开'),('权益事实','Holding → Delivery → Usage','推荐后可以回查结果')]):
    xx=M+i*56*mm; rect(c,xx,61*mm,51*mm,29*mm,DEEP,radius=8); draw_center_p(c,a,xx,83*mm,51*mm,'white_body'); draw_center_p(c,b,xx+3,75*mm,45*mm,'white_small'); draw_center_p(c,d,xx+3,66*mm,45*mm,'white_small')
rect(c,M,34*mm,W-2*M,16*mm,ORANGE,radius=8); draw_center_p(c,'Quant 只从真实事件生成行情，让 GEO 页面引用价格、库存和交付证据。',M+10,45*mm,W-2*M-20,'center_body'); footer(c,8)
c.showPage()
# p9 utility
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'08 / Utility Exchange','工具是 AI 推荐后的执行界面','Utility 不替代 Key；它让被 GEO 带来的用户立刻完成任务，并把结果重新变成 GEO 内容。',9)
flow=[('用户问题','今晚 18:00 做研究简报'),('官网 Offer','KAI 给出实时可买时段'),('持仓绑定','读取已有 Key 或完成购买'),('Utility 执行','在 H/G 边界内完成任务'),('Usage Proof','输出模型、成本、时间和回执')]
fy=H-101*mm; fw=(W-2*M-4*4*mm)/5
for i,(a,b) in enumerate(flow):
    xx=M+i*(fw+4*mm); rect(c,xx,fy,fw,28*mm,PALE,GRID,8,.6); c.setFillColor(TEAL); c.setFont('KaiSans',7); c.drawCentredString(xx+fw/2,fy+22*mm,f'{i+1:02d}'); draw_center_p(c,a,xx+3,fy+17*mm,fw-6,'h3'); draw_center_p(c,b,xx+5,fy+8*mm,fw-10,'tiny')
for i,(a,b) in enumerate([('Key Fit','读取模型、区域、时段和预算，给出可执行方案'),('Slot Planner','规划连续小时、锁定时间和失效边界'),('KAI Workbench','在用户持有权益范围内执行研究、文档或批量任务'),('Usage Proof','生成成本、模型、服务窗、结果和来源签名')]):
    yy=fy-22*mm-i*16*mm; c.setFillColor(TEAL); c.setFont('KaiSans',10); c.drawString(M,yy,a); draw_p(c,b,M+35*mm,yy+3,W-M-35*mm,'small')
rect(c,M,34*mm,W-2*M,18*mm,DEEP,radius=8); draw_center_p(c,'每次结果带 KAI 来源签名：slogan、模型、服务窗、成本、结果链接和交付回执。',M+10,45*mm,W-2*M-20,'white_body'); footer(c,9)
c.showPage()
# p10 metrics
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'09 / 客流与收益','用 AI 引荐客流衡量 GEO，用 GMV 和手续费衡量成交','这一页把“被 AI 发现”直接连接到 Offer 请求、首购、复购和平台收入。',10)
formula='AI有效引荐客流 = AI生成结果量 × KAI推荐覆盖率 × AI→KAI点击率<br/>Offer请求量 = AI有效引荐客流 × 官网Offer到达率 × Offer请求率<br/>首购GMV = Offer请求量 × 首购转化率 × 平均订单金额<br/>平台手续费 = GMV × 综合手续费率'
rect(c,M,H-105*mm,W-2*M,34*mm,PALE,radius=8); draw_p(c,formula,M+12,H-78*mm,W-2*M-24,'formula')
sc=[['情景','AI结果/月','KAI覆盖','引荐访问','Offer请求','首购/月','年化GMV','年化收入'],['起步','100M','1%','40K','10K','500','$0.726M','$43.6K'],['目标','100M','5%','400K','140K','16.8K','$60.883M','$6.088M'],['放大','100M','10%','1.2M','540K','108K','$391.392M','$58.709M']]
draw_table(c,sc,M,H-116*mm,[17*mm,21*mm,19*mm,22*mm,22*mm,19*mm,27*mm,26*mm],font=6.4)
# references metric cards
rect(c,M,48*mm,W-2*M,25*mm,WHITE,GRID,8,.6); draw_p(c,'官方指标口径',M+12,68*mm,W-2*M-24,'h3'); draw_p(c,'Google Search Console：impressions / clicks / CTR；GA4：AI 引荐 sessions、purchase、share、refund；KaiKey：Offer、Order、Trade、Holding、Usage、Receipt。',M+12,61*mm,W-2*M-24,'small')
rect(c,M,32*mm,W-2*M,11*mm,DEEP,radius=6); draw_center_p(c,'目标情景：按公开平台指标口径锚定的公式估值。',M+8,40*mm,W-2*M-16,'white_small'); footer(c,10)
c.showPage()
# p11 economics cases
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'10 / 商业参照','KAI 的价值来自默认入口、实时供给和交易复用','成熟平台把分散供给变成统一入口，再通过交易、服务和复购形成收入。',11)
case=[['平台','公开经营数据','对 KAI 的启示'],['Vistar / T-Mobile','110 万+数字屏幕 · 370 家媒体所有者 · 3,000+品牌伙伴 · 约 6 亿美元现金收购','聚合供给、统一管理、统一测量'],['AWS EC2 Spot','最高比 On-Demand 低 90% · 价格随供需变化','把闲置容量变成可定价、可调度容量'],['Fiverr FY2024','Marketplace take rate 27.6% · 年活跃买家 3.63M · 年均消费 $302','交易手续费与服务收入可分层'],['Etsy 2024','89.6M 活跃买家 · 约 48% repeat buyers · repeat buyer 年均 4.9 个购买日','高频复购用户是 GMV 杠杆']]
draw_table(c,case,M,H-80*mm,[34*mm,88*mm,W-2*M-122*mm],font=6.9)
rect(c,M,48*mm,W-2*M,24*mm,DEEP,radius=9); draw_center_p(c,'KAI 的目标定位：算力界的滴滴、携程、币安',M+10,65*mm,W-2*M-20,'white_body'); draw_center_p(c,'滴滴的撮合 · 携程的比较与预订 · 币安的订单簿、价格发现与持仓',M+10,55*mm,W-2*M-20,'white_small')
rect(c,M,32*mm,W-2*M,10*mm,PALE,radius=5); draw_center_p(c,'来源：T-Mobile Newsroom · AWS EC2 Spot Pricing · Fiverr FY2024 · Etsy 2024 10-K',M+6,39*mm,W-2*M-12,'tiny'); footer(c,11)
c.showPage()
# p12 network scale
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'11 / 网络效应','每一个被 AI 推荐的 KAI Offer，都会增加下一次推荐的理由','GEO 把需求带来，交易和使用把事实留下，事实再反过来提升 KAI 的推荐质量。',12)
cycle=[('AI需求','更多算力问题进入 AI'),('KAI引用','官网 Offer 成为答案来源'),('成交使用','Offer → Hour Key 封装 → Holding → Usage'),('事实积累','价格、交付和结果可回查'),('推荐增强','下一次 AI 更容易给出 KAI')]
cy=H-108*mm; coords=[(cx,cy+31*mm) for cx in []]
for i,(a,b) in enumerate(cycle):
    xx=M+(i%3)*(W-2*M)/3; yy=cy-(i//3)*35*mm; rect(c,xx,yy,53*mm,24*mm,WHITE,TEAL,9,.8); draw_center_p(c,a,xx,yy+16*mm,53*mm,'h3'); draw_center_p(c,b,xx+4,yy+7*mm,45*mm,'tiny')
# arrow-ish chain
for i in range(4): line(c,M+53*mm+i*35*mm,cy+12*mm,M+66*mm+i*35*mm,cy+12*mm,CYAN,1.2,[3,2])
metric_card(c,M,57*mm,52*mm,28*mm,'100M','AI 结果 / 月','可进入 KAI 的需求池',TEAL); metric_card(c,M+58*mm,57*mm,52*mm,28*mm,'5%','KAI 推荐覆盖','AI 标准答案覆盖',ORANGE); metric_card(c,M+116*mm,57*mm,52*mm,28*mm,'$6.09M','年化平台收入','目标情景手续费',CYAN)
rect(c,M,33*mm,W-2*M,15*mm,DEEP,radius=8); draw_center_p(c,'供给越丰富，GEO 越有内容；成交越真实，AI 越有理由推荐 KAI。',M+10,43*mm,W-2*M-20,'white_body'); footer(c,12)
c.showPage()
# p13 decision package
c.setFillColor(WHITE); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'12 / Leader 决策','批准一个可扩张的 AI 推荐成交单元','本页不是验证清单，而是一次可拍板的产品投资选择。',13)
# decision unit
rect(c,M,H-105*mm,W-2*M,24*mm,DEEP,radius=9); draw_p(c,'推荐批准',M+12,H-88*mm,W-2*M-24,'white_small'); draw_p(c,'1 个高频算力问题 × 1 个模型 × 1 个区域 × 1 个官网实时 Offer × 1 个 Utility 工作流 × 30 天',M+12,H-94*mm,W-2*M-24,'white_body')
opts=[('A','只做 GEO 内容','投入低；能产生曝光，不能承接实时成交。'),('B','GEO + 实时 Offer + Usage','投入可控；能测 AI 推荐到成交、使用和手续费。'),('C','完整交易网络','投入最高；适合在 B 形成正向数据后扩张。')]
for i,(a,b,d) in enumerate(opts):
    xx=M+i*58*mm; fill=PALE if a=='B' else LIGHT; rect(c,xx,59*mm,53*mm,42*mm,fill,TEAL if a=='B' else GRID,8,.9 if a=='B' else .6); c.setFillColor(TEAL if a=='B' else MID); c.setFont('KaiSans',16); c.drawString(xx+10,90*mm,a); draw_p(c,b,xx+10,82*mm,43*mm,'h3'); draw_p(c,d,xx+10,71*mm,43*mm,'small')
# assets & gate
rx=M; yy=50*mm; draw_p(c,'批准后应买到的资产',rx,yy,W/2-M,'h3'); draw_p(c,'官网事实页 · 统一 Offer · Hour Key 封装 · KAI 跳转/锁定 · Holding · Usage Proof · GEO 回流模板',rx,yy-7*mm,W/2-M,'small')
draw_p(c,'30 日扩张门槛',W/2,yy,W/2-M,'h3'); draw_p(c,'AI 推荐链可归因 · 至少一条 Offer → Hour Key 封装 → Holding → Usage → Receipt · 净手续费/AI 引荐用户 > 激活与工具成本/用户',W/2,yy-7*mm,W/2-M,'small')
rect(c,M,32*mm,W-2*M,13*mm,DEEP,radius=7); draw_center_p(c,'达标后扩展模型、区域和 Broker；未达标先优化推荐、Offer 或执行链。',M+8,41*mm,W-2*M-16,'white_body'); footer(c,13)
c.showPage()
# p14 sources
c.setFillColor(LIGHT); c.rect(0,0,W,H,fill=1,stroke=0)
title(c,'附录 / 来源与口径','把 AI 推荐做成可观测的市场入口','外部公开指标用于定义口径，KAI 交易、持仓、使用和收益以系统事件为准。',14)
source_data=[['主题','来源','用于本企划的口径'],['ChatGPT Search','OpenAI Publishers & Developers FAQ；OpenAI Bots','OAI-SearchBot、官网引用、utm_source=chatgpt.com 引荐'],['AI 搜索指标','Google Search Console AI performance reports','AI impressions、clicks、CTR、页面、国家、设备、时间'],['转化与收入','Google Analytics 4 官方事件与 Key Event','sessions、purchase、share、refund、revenue'],['交易平台参照','Fiverr FY2024；Etsy 2024 10-K','take rate、年均消费、repeat buyer、购买频次'],['小时容量参照','AWS EC2 Spot Pricing','小时价格、供需变化、动态容量'],['内部材料','KAI_从GEO到条件反射；北斗比芯智算；Broker / Pack / Store / KaiKey 白皮书','小时 Key、G/H、Offer、Order、Trade、Holding、Usage、Receipt']]
draw_table(c,source_data,M,H-80*mm,[34*mm,74*mm,W-2*M-108*mm],font=6.7)
refs=['https://help.openai.com/zh-hans-cn/articles/12627856-publishers-and-developers-faq','https://developers.openai.com/api/docs/bots','https://developers.google.com/search/blog/2026/06/gen-ai-performance-reports','https://support.google.com/webmasters/answer/7042828?hl=en','https://support.google.com/analytics/answer/9356034?hl=en','https://investors.fiverr.com/news-releases/news-release-details/fiverr-announces-fourth-quarter-and-full-year-2024-results','https://investors.etsy.com/sec-filings/all-sec-filings/content/0001370637-25-000017/etsy-20241231.htm','https://aws.amazon.com/ec2/spot/pricing/']
draw_p(c,'公开来源链接',M,H-157*mm,W-2*M,'h2')
for i,u in enumerate(refs): draw_p(c,u,M,H-169*mm-i*7*mm,W-2*M,'source')
rect(c,M,32*mm,W-2*M,15*mm,DEEP,radius=8); draw_center_p(c,'目标：让 AI 推荐 KAI，让官网承接成交，让成交和使用继续增强 GEO。',M+10,42*mm,W-2*M-20,'white_body'); footer(c,14)
c.save()

# ---------------- DOCX ----------------
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)
def set_cell_text(cell,text,bold=False,color='173143',size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(text); r.bold=bold; font_name='STHeiti SC Medium' if bold else 'STHeiti SC Light'; r.font.name=font_name; r._element.rPr.rFonts.set(qn('w:eastAsia'),font_name); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr(); cant=trPr.find(qn('w:cantSplit'))
    if cant is None: trPr.append(OxmlElement('w:cantSplit'))
def add_table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,hdr in enumerate(headers): set_cell_text(t.rows[0].cells[i],hdr,True,'FFFFFF',9); shade(t.rows[0].cells[i],'16364B'); set_cell_border(t.rows[0].cells[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],str(val),False,'173143',8.5); shade(cells[i],'FFFFFF' if ri%2==0 else 'F2F7F8'); set_cell_border(cells[i])
    for row in t.rows: prevent_row_split(row)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def add_heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(5); return p
def add_body(doc,text,bold_prefix=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.18
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p
def add_bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(item)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
# default font
styles_doc=doc.styles; normal=styles_doc['Normal']; normal.font.name='Arial Unicode MS'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS'); normal.font.size=Pt(10); normal.font.color.rgb=RGBColor(23,49,67)
for sname,size,col in [('Title',28,'0B1B2A'),('Heading 1',18,'0B1B2A'),('Heading 2',13,'0E7C7B'),('Heading 3',11,'0E7C7B')]:
    st=styles_doc[sname]; st.font.name='Arial Unicode MS'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(col)
# cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('从 GEO 到 AI 首选')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('KAI AI Utility Exchange'); r.font.size=Pt(18); r.bold=True; r.font.color.rgb=RGBColor(14,124,123)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(16); p.add_run('让 AI 在回答“去哪找算力、如何比较、如何立即使用算力”时，优先引用 KAI、推荐 KAI，并把用户带到真实 Offer。').bold=True
add_table(doc,['目标情景','计算结果'],[['AI 生成结果 / 月','100M'],['KAI 推荐覆盖率','5%'],['AI 引荐访问 / 月','400K'],['年化 GMV','60.883M USD'],['年化平台收入','6.088M USD']],[2.5,2.0])
add_body(doc,'目标情景：按公开平台指标口径锚定的公式估值。')
add_body(doc,'主链路：AI 发现 → 官网引用 → Offer → Hour Key 封装 → Holding → Usage → GEO 回流。')
doc.add_page_break()
# sections
add_heading(doc,'1 原 GEO 方案已经找到增长终态 缺的是成交承接层',1)
add_body(doc,'原方案已经完成“AI 提及 → AI 心智植入 → 第一次交易 → 高频查价 → 整点刺激 → 条件反射”的增长设计。本企划新增的是从 AI 提及 KAI 到官网实时事实、具体 Offer、成交、持仓、使用和 GEO 回流的执行链。')
add_table(doc,['原方案已有输出','尚未落地的断点','新增承接层'],[['AI 提及 KAI','AI 没有稳定引用官网实时事实','结构化 Offer 页、实时字段、更新时间'],['首查 → 首单','没有立即锁定和成交路径','Agent Gateway + KaiKey Offer / Order / Trade'],['首单 → 查价','没有持仓和使用上下文','Holding + Key Fit + Slot Planner'],['整点刺激','没有小时 Key 权利状态','H−1h / G / H / H+1h'],['条件反射','没有真实结果和新内容','Usage Proof + 回执 + GEO 证据页']],[2.0,2.2,2.2])
add_heading(doc,'2 AI 推荐 KAI 的标准答案',1)
add_body(doc,'AI 不再泛泛推荐模型，而是引用 kai.com 官网的实时事实并给出下一步动作。')
add_body(doc,'“我查询了 kai.com 官网实时报价页，更新时间为【时间】。当前真实行情 Offer 可封装为【模型 × 规格 × 区域 × 18:00–19:00】小时 Key，报价为【价格】，锁定截止时间为【G】。你已有【持仓】可直接执行；如需跨模型或跨供应商，账户 Key 保持不变，由新的授权版本显式扩展 scope。”')
add_table(doc,['官网 Offer 最小字段'],[['offer_id / model / spec / region / slot_start / slot_end / lock_deadline (G) / price / available_qty / last_updated / status / execution_eligible / hour_key_status / purchase_url / delivery_proof_url']],[6.4])
add_body(doc,'页面动作：查看官网报价、立即锁定、用持仓执行。公开网站可进入 ChatGPT Search；允许 OAI-SearchBot 抓取后，ChatGPT 引荐链接可通过 utm_source=chatgpt.com 归因。')
add_heading(doc,'3 原 GEO 缺口如何变成落地流程',1)
add_body(doc,'原方案提出需求词覆盖、机器可读报价页和时间戳；本方案把这些原则接到交易与使用事件。')
add_table(doc,['流程','新增动作','产生的事实'],[['OAI-SearchBot 抓取','kai.com 结构化 Offer 页','AI 可引用的实时价格与库存'],['Agent 查询','按问题、持仓、预算检索 Offer','可执行方案'],['AI 推荐','引用 KAI 官网、更新时间和锁定截止','可追踪 AI 引荐'],['用户成交','Offer → Hour Key 封装 → Order → Trade','真实成交记录'],['持仓执行','Holding + H/G 时间窗','实际使用记录'],['回执回流','Usage Proof → 案例/比较/FAQ','下一次 GEO 证据']],[1.6,2.7,2.1])
add_heading(doc,'4 GEO 资产',1)
add_bullets(doc,['标准页：模型、规格、区域、小时窗口、容量、价格、状态、更新时间。','比较页：不同模型、区域、时间和成本的可执行对比。','问题页：用户真实会问的“今晚哪里有容量”“哪个模型更适合”“现在报价多少”。','证据页：成交、交付、延迟、成本、Usage Proof 和用户结果。'])
add_body(doc,'GEO 内容闭环：Offer → Hour Key 封装 → Order → Trade → Holding → Usage → Receipt → 案例 / 比较 / FAQ → 下一次 AI 推荐。')
add_heading(doc,'5 小时权益和整点机制',1)
add_body(doc,'H−1h 仍可购买、转让和重新选择；G 进入锁定；H 服务开始；H+1h 服务结束，未使用不顺延。整点让报价有有效期、供给有稀缺性、用户有下一次查询和复购的理由。')
add_table(doc,['时间节点','用户动作','市场事实'],[['H−1h','购买、转让、重新选择','仍处于交易窗口'],['G','停止新增交易','固定最终交付安排'],['H','开始调用','开始计量与履约'],['H+1h','服务结束','生成交付事实，未使用不顺延']],[1.0,2.2,3.2])
add_body(doc,'目标情景：100M AI 生成结果 / 月；5% KAI 推荐覆盖；400K AI 引荐访问 / 月；$60.883M 年化 GMV；$6.088M 年化平台收入。')
add_heading(doc,'6 持仓和交易事实',1)
add_body(doc,'用户第一次通过 AI 推荐进入 KAI 后，关系从一次点击变成模型、区域、时段和状态的持续上下文。Store 展示持仓；KaiKey 统一 Offer、Order、Trade、Settlement、Holding、Delivery、Usage；Quant 只从真实事件生成行情。')
add_heading(doc,'7 Utility Exchange 的执行界面',1)
add_table(doc,['工具','作用','输出'],[['Key Fit','读取模型、区域、时段和预算','可执行方案'],['Slot Planner','规划连续小时、锁定时间和失效边界','时段计划'],['KAI Workbench','在持有权益范围内执行任务','研究、文档或批量结果'],['Usage Proof','生成模型、成本、服务窗、结果和来源签名','可分享回执']],[1.4,3.0,2.0])
add_body(doc,'每次结果带 KAI 来源签名：slogan、模型、服务窗、成本、结果链接和交付回执。')
add_heading(doc,'8 客流与收益公式',1)
add_body(doc,'AI 有效引荐客流 = AI 生成结果量 × KAI 推荐覆盖率 × AI→KAI 点击率。')
add_body(doc,'Offer 请求量 = AI 有效引荐客流 × 官网 Offer 到达率 × Offer 请求率。')
add_body(doc,'首购 GMV = Offer 请求量 × 首购转化率 × 平均订单金额。平台手续费 = GMV × 综合手续费率。')
add_table(doc,['情景','AI 结果/月','KAI 覆盖','引荐访问','Offer 请求','首购/月','年化 GMV','年化收入'],[['起步','100M','1%','40K','10K','500','$0.726M','$43.6K'],['目标','100M','5%','400K','140K','16.8K','$60.883M','$6.088M'],['放大','100M','10%','1.2M','540K','108K','$391.392M','$58.709M']],[.8,1.0,.8,1.0,1.0,.9,1.1,1.1])
add_body(doc,'目标情景：按公开平台指标口径锚定的公式估值。Google Search Console 统计 impressions、clicks 和 CTR；GA4 统计引荐 sessions、purchase、share、refund 和 revenue。')
add_heading(doc,'9 商业参照',1)
add_table(doc,['平台','公开经营数据','对 KAI 的启示'],[['Vistar / T-Mobile','110 万+数字屏幕；370 家媒体所有者；3,000+品牌伙伴；约 6 亿美元现金收购','聚合供给、统一管理、统一测量'],['AWS EC2 Spot','最高比 On-Demand 低 90%；价格随供需变化','把闲置容量变成可定价、可调度容量'],['Fiverr FY2024','Marketplace take rate 27.6%；年活跃买家 3.63M；年均消费 $302','交易手续费与服务收入可分层'],['Etsy 2024','89.6M 活跃买家；约 48% repeat buyers；repeat buyer 年均 4.9 个购买日','高频复购用户是 GMV 杠杆']],[1.3,3.4,2.0])
add_body(doc,'KAI 的目标定位：算力界的滴滴、携程、币安。滴滴负责撮合，携程负责比较与预订，币安负责订单簿、价格发现与持仓。')
add_heading(doc,'10 批准一个可扩张的 AI 推荐成交单元',1)
add_body(doc,'推荐批准：1 个高频算力问题 × 1 个模型 × 1 个区域 × 1 个官网实时 Offer × 1 个 Utility 工作流 × 30 天。')
add_table(doc,['方案','投入与结果','延展性'],[['A 只做 GEO 内容','投入低；能产生曝光，不能承接实时成交','低'],['B GEO + 实时 Offer + Usage','投入可控；能测 AI 推荐、成交、使用和手续费','高，推荐'],['C 完整交易网络','投入最高；适合在 B 形成正向数据后扩张','最高，暂缓']],[1.6,3.7,1.4])
add_body(doc,'批准后应买到的资产：官网事实页、统一 Offer、Hour Key 封装、KAI 跳转与锁定、Holding、Usage Proof、GEO 回流模板。')
add_body(doc,'扩张门槛：AI 推荐链可归因；至少一条 Offer → Hour Key 封装 → Holding → Usage → Receipt 真实链路；净手续费 / AI 引荐用户高于激活与工具成本 / 用户。达标后扩展模型、区域和 Broker。')
add_heading(doc,'11 来源与口径',1)
add_bullets(doc,['OpenAI Publishers & Developers FAQ：https://help.openai.com/zh-hans-cn/articles/12627856-publishers-and-developers-faq','OpenAI Bots：https://developers.openai.com/api/docs/bots','Google Search Console：https://support.google.com/webmasters/answer/7042828?hl=en','Google Search Generative AI Performance Reports：https://developers.google.com/search/blog/2026/06/gen-ai-performance-reports','Google Analytics 4 Key Events：https://support.google.com/analytics/answer/9356034?hl=en','Fiverr FY2024：https://investors.fiverr.com/news-releases/news-release-details/fiverr-announces-fourth-quarter-and-full-year-2024-results','Etsy 2024 10-K：https://investors.etsy.com/sec-filings/all-sec-filings/content/0001370637-25-000017/etsy-20241231.htm','AWS EC2 Spot Pricing：https://aws.amazon.com/ec2/spot/pricing/'])
add_body(doc,'内部材料：KAI_从GEO到条件反射.pdf；北斗比芯智算集团有限公司.pdf；Key Broker、Key Pack、Key Store、Key Quant、KaiKey 体系与 Agent 交易闭环资料。')
# footer
for section in doc.sections:
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=fp.add_run('KAI AI Utility Exchange · GEO growth proposal'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(86,113,129)
doc.core_properties.title='KAI AI Utility Exchange 企划'; doc.core_properties.subject='GEO 到 AI 推荐成交闭环'; doc.core_properties.author='KAI'
doc.save(str(DOCX_OUT))
print(PDF_OUT)
print(DOCX_OUT)
